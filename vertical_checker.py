#!/usr/bin/env python3
"""
Vertical Grid Checker — runner, portal-config lint, profile validation, Excel report.

  python vertical_checker.py --csv final.csv --profile profiles/banca_motor.json \
      [--source GRID.xlsx] [--rto rto.xlsx] [--reference last_month.csv] [--config portal_config.json] [--report r.xlsx]

Exit code 0 = no ERROR, 1 = errors, 2 = checker failed.
"""
import argparse, json, os, re, sys, traceback
from collections import defaultdict
import pandas as pd
from vc_core import ERROR, WARN, INFO, OPS, NO_VALUE_OPS, Findings, read_source, canon, to_num, describe_conds
from vc_checks import CHECK_TYPES, run_check
from vc_reconcile import detect_bands, reconcile


# ── portal config lint ────────────────────────────────────────────────────────
def lint_config(cfg, prof, T, F, bands):
    C = 'CONFIG'
    std = cfg.get('std', cfg)
    src = prof.get('source') or {}
    try: hr = int(std.get('header_row') or 0)
    except ValueError: hr = 0
    if src.get('header_row') and hr and hr != int(src['header_row']):
        F.add(C, 'L01', 'Portal header row differs from the sheet layout', ERROR, 1, expected=src['header_row'], actual=hr,
              fix='With the wrong header row the portal reads the band hint row (0-1, 1-6…) and produces no detriff expansion.')
    if src.get('sheet') and std.get('sheet_name') and std['sheet_name'] != src['sheet']:
        F.add(C, 'L02', 'Portal sheet differs from the profile sheet', WARN, 1, expected=src['sheet'], actual=std['sheet_name'])
    th = [x.strip() for x in str(std.get('target_hdrs') or '').split('\n') if x.strip()] if isinstance(std.get('target_hdrs'), str) else (std.get('target_headers') or [])
    exp = prof.get('expected_columns') or []
    if th and exp:
        miss = [c for c in exp if c not in th]; extra = [c for c in th if c not in exp]
        if miss: F.add(C, 'L03', 'Target headers missing template columns', ERROR, len(miss), actual=', '.join(miss))
        if extra: F.add(C, 'L03b', 'Target headers not in template', WARN, len(extra), actual=', '.join(extra))
    mp = std.get('map_json') or std.get('column_mapping') or {}
    if isinstance(mp, str):
        try: mp = json.loads(mp)
        except ValueError: F.add(C, 'L04', 'Column mapping JSON is invalid', ERROR, 1); mp = {}
    if T is not None:
        for tc, sc in mp.items():
            if isinstance(sc, str) and not sc.startswith('__') and sc not in T.columns:
                F.add(C, 'L05', 'Mapping refers to a source column that does not exist — the portal writes the column NAME as the value', ERROR, 1,
                      column=tc, actual=sc, fix='Rename the source column in the mapping, or use __LITERAL__:value.')
    by_src = defaultdict(list)
    for tc, sc in mp.items():
        if isinstance(sc, str) and not sc.startswith('__'): by_src[sc].append(tc)
    for sc, tcs in by_src.items():
        ll = [t for t in tcs if re.search(r'\bLl\*?$', t)]; ul = [t for t in tcs if re.search(r'\bUl\*?$', t)]
        if ll and ul:
            F.add(C, 'L06', 'LL and UL mapped from the same source column (needs a manual split)', WARN, 1, column=', '.join(ll + ul), actual=sc,
                  fix="e.g. 'Vehicle Age' = '< 5' must become LL/UL numbers; add a source rule to verify the split.")
    rate_out = ((prof.get('reconcile') or {}).get('rate') or {}).get('output')
    rate_map = [tc for tc, sc in mp.items() if sc == '__DETRIFF_VALUE__']
    if rate_out and mp and rate_out not in rate_map:
        F.add(C, 'L07', 'Rate column in portal mapping differs from the profile', ERROR, 1, expected=rate_out, actual=', '.join(rate_map) or '(none)')
    scale = float(((prof.get('reconcile') or {}).get('rate') or {}).get('scale', 1) or 1)
    if scale != 1 and rate_out:
        txs = std.get('transformations') or []
        has = any(t.get('column') == rate_out and any(o.get('type') in ('math_op',) for o in t.get('ops', [])) for t in txs)
        if not has:
            F.add(C, 'L08', f'No math transform on the rate column but the profile expects ×{scale:g}', WARN, 1, column=rate_out,
                  fix='Add a math_op (x*100) transformation in Step 4, or set the rate scale in the profile to 1.')


# ── profile validation ────────────────────────────────────────────────────────
def _vconds(conds, where, errs):
    for i, c in enumerate(conds or [], 1):
        if c.get('op') not in OPS: errs.append(f'{where}: condition {i} has unknown op "{c.get("op")}"')
        if not c.get('col'): errs.append(f'{where}: condition {i} has no column')
        if c.get('op') in ('regex', 'not_regex'):
            try: re.compile(str(c.get('value', '')))
            except re.error as e: errs.append(f'{where}: condition {i} bad regex ({e})')


def validate_profile(p):
    errs, warns = [], []
    if not isinstance(p, dict): return ['Profile must be a JSON object'], []
    if not p.get('expected_columns'): warns.append('expected_columns is empty — template checks will not run')
    ids = set()
    for i, c in enumerate(p.get('checks') or [], 1):
        w = f'Check {c.get("id") or i}'
        if c.get('id') in ids: errs.append(f'{w}: duplicate id')
        ids.add(c.get('id'))
        if c.get('type') not in CHECK_TYPES: errs.append(f'{w}: unknown type "{c.get("type")}"'); continue
        if c.get('target', 'output') not in CHECK_TYPES[c['type']]['targets']:
            errs.append(f'{w}: type {c["type"]} cannot run on {c.get("target")}')
        if c.get('severity', ERROR) not in (ERROR, WARN, INFO): errs.append(f'{w}: bad severity')
        _vconds(c.get('where'), w + ' where', errs)
        if c['type'] == 'rule':
            _vconds((c.get('params') or {}).get('if'), w + ' IF', errs); _vconds((c.get('params') or {}).get('then'), w + ' THEN', errs)
            if not (c.get('params') or {}).get('then'): errs.append(f'{w}: rule has no THEN condition')
        if c['type'] == 'pattern':
            try: re.compile((c.get('params') or {}).get('regex', ''))
            except re.error as e: errs.append(f'{w}: bad regex ({e})')
    for i, r in enumerate(p.get('source_rules') or [], 1):
        w = f'Source rule {r.get("id") or i}'
        if r.get('id') in ids: errs.append(f'{w}: duplicate id')
        ids.add(r.get('id'))
        if not r.get('when'): errs.append(f'{w}: needs at least one WHEN condition')
        _vconds(r.get('when'), w + ' WHEN', errs)
        for j, g in enumerate(r.get('expect') or [], 1): _vconds(g.get('conds'), f'{w} expectation {j}', errs)
        if not (r.get('expect') or (r.get('effects') or {}).get('no_output') or (r.get('effects') or {}).get('cap_band_ul')):
            warns.append(f'{w}: no expectation or effect — it only reports coverage')
    rc = p.get('reconcile') or {}
    if rc.get('enabled'):
        if not rc.get('keys'): errs.append('Reconcile is enabled but has no keys')
        for k, v in (rc.get('tokens') or {}).items():
            if (v or {}).get('policy') not in ('expect', 'optional', 'skip'): errs.append(f'Token "{k}": policy must be expect / optional / skip')
        if not (p.get('source') or {}).get('band_map'): errs.append('Reconcile is enabled but source.band_map is empty')
        for k, v in ((p.get('source') or {}).get('band_map') or {}).items():
            if not (isinstance(v, list) and len(v) == 2 and to_num(v[0]) is not None and to_num(v[1]) is not None):
                errs.append(f'Band "{k}" must be [LL, UL]')
    return errs, warns


# ── report ────────────────────────────────────────────────────────────────────
def write_report(path, F):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    base = Font(name='Arial', size=10); bold = Font(name='Arial', size=10, bold=True)
    hf = Font(name='Arial', size=10, bold=True, color='FFFFFF'); hfill = PatternFill('solid', start_color='1F2A44')
    sf = {ERROR: PatternFill('solid', start_color='F8CBAD'), WARN: PatternFill('solid', start_color='FFE699'), INFO: PatternFill('solid', start_color='DDEBF7')}
    ws = wb.active; ws.title = 'Summary'
    ws['A1'] = 'Vertical Grid Checker Report'; ws['A1'].font = Font(name='Arial', size=14, bold=True)
    r = 3
    for k, v in F.meta.items():
        ws.cell(r, 1, k).font = bold; ws.cell(r, 2, str(v)).font = base; r += 1
    e = sum(i['severity'] == ERROR for i in F.items); w = sum(i['severity'] == WARN for i in F.items)
    ws.cell(r, 1, 'Verdict').font = bold
    ws.cell(r, 2, 'FAIL — do not upload' if e else ('PASS with warnings' if w else 'PASS')).font = Font(name='Arial', size=12, bold=True, color='C00000' if e else '375623')
    r += 1; ws.cell(r, 1, 'Findings').font = bold; ws.cell(r, 2, f'{e} errors, {w} warnings').font = base
    ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 90
    order = {ERROR: 0, WARN: 1, INFO: 2}
    wd = wb.create_sheet('Details')
    cols = ['severity', 'layer', 'check_id', 'check', 'column', 'expected', 'actual', 'count', 'examples', 'fix', 'rule']
    for j, h in enumerate(cols, 1):
        c = wd.cell(1, j, h.replace('_', ' ').title()); c.font = hf; c.fill = hfill
    for i, it in enumerate(sorted(F.items, key=lambda x: (order[x['severity']], x['layer'], x['check_id'])), 2):
        for j, h in enumerate(cols, 1):
            c = wd.cell(i, j, it[h]); c.font = base; c.alignment = Alignment(wrap_text=h in ('check', 'examples', 'fix', 'actual', 'expected', 'rule'), vertical='top')
        wd.cell(i, 1).fill = sf[it['severity']]
    for col, wdt in zip('ABCDEFGHIJK', (9, 9, 9, 46, 26, 28, 40, 9, 50, 50, 36)): wd.column_dimensions[col].width = wdt
    wd.freeze_panes = 'A2'; wd.auto_filter.ref = f'A1:K{len(F.items) + 1}'
    if F.samples:
        wm = wb.create_sheet('Mismatch samples')
        mc = ['source_cell', 'csv_line', 'column', 'expected', 'actual']
        for j, h in enumerate(mc, 1):
            c = wm.cell(1, j, h.replace('_', ' ').title()); c.font = hf; c.fill = hfill
        for i, m in enumerate(F.samples[:3000], 2):
            for j, h in enumerate(mc, 1): wm.cell(i, j, m[h]).font = base
        for col, wdt in zip('ABCDE', (40, 10, 30, 24, 24)): wm.column_dimensions[col].width = wdt
    wb.save(path)


# ── runner ────────────────────────────────────────────────────────────────────
def run(csv, profile, source=None, rto=None, reference=None, config=None, report=None, label=None, quiet=False):
    prof = profile if isinstance(profile, dict) else json.load(open(profile))
    F = Findings()
    try:
        out = pd.read_csv(csv, dtype=str, keep_default_na=False)
    except pd.errors.EmptyDataError:
        out = pd.DataFrame()
    out.columns = [str(c).strip() for c in out.columns]
    if out.shape[1] == 0:
        F.add('OUTPUT', 'O00', 'The output file has no columns (empty CSV)', ERROR, 1,
              fix='In the portal, open Step 3 (Column Mapping) before processing so the mapping is built, then process again.')
        F.meta = {'Output file': label or os.path.basename(csv), 'Rows': '0', 'Profile': prof.get('name', '')}
        if report: write_report(report, F)
        return F
    T = hdr = letters = None; bands = []
    sheet = (prof.get('source') or {}).get('sheet', '')
    if source:
        try:
            T, hdr, letters = read_source(source, prof.get('source') or {})
            bands = detect_bands(T, (prof.get('source') or {}).get('band_map'))
        except Exception as ex:
            F.add('SOURCE', 'R00', f'Could not read the source sheet: {ex}', ERROR, 1)
    ref_df = pd.read_csv(reference, dtype=str, keep_default_na=False) if reference else None
    rto_map = _rto(rto) if rto else None
    ctx = {'hdr': hdr or {}, 'rto': rto_map, 'reference': ref_df, 'bands': bands}
    skipped_src = 0
    for chk in prof.get('checks') or []:
        if not chk.get('enabled', True): continue
        tgt = chk.get('target', 'output')
        df = out if tgt == 'output' else T
        if df is None: skipped_src += 1; continue
        try: run_check(chk, df, F, prof, ctx)
        except Exception as ex:
            F.add('CHECKER', chk.get('id', '?'), f'Check could not run: {ex}', WARN, 1, fix='Fix this check’s parameters in the profile.')
    if skipped_src:
        F.add('SOURCE', 'R00i', f'{skipped_src} source checks skipped (no source workbook given)', INFO, 0)
    if T is not None and (prof.get('reconcile') or {}).get('enabled'):
        try: reconcile(T, sheet, letters, hdr, bands, out, prof, F, ctx)
        except Exception as ex:
            F.add('CHECKER', 'R', f'Reconciliation failed: {ex}', ERROR, 1, fix=traceback.format_exc()[-300:])
    if config:
        try: lint_config(config if isinstance(config, dict) else json.load(open(config)), prof, T, F, bands)
        except Exception as ex: F.add('CHECKER', 'L', f'Config lint failed: {ex}', WARN, 1)
    F.meta = {'Output file': label or os.path.basename(csv), 'Rows': f'{len(out):,}', 'Profile': prof.get('name', ''),
              'Source': f"{os.path.basename(source)} :: {sheet}" if source else '(not given)',
              'Header cells': ', '.join(f'{k}={v}' for k, v in (hdr or {}).items()) or '—'}
    if report: write_report(report, F)
    if not quiet:
        e = [i for i in F.items if i['severity'] == ERROR]; w = [i for i in F.items if i['severity'] == WARN]
        print(f"\n{'=' * 90}\n{F.meta['Output file']} — {len(out):,} rows — {len(e)} errors / {len(w)} warnings\n{'=' * 90}")
        for it in sorted(F.items, key=lambda x: ({ERROR: 0, WARN: 1, INFO: 2}[x['severity']], x['layer'], x['check_id'])):
            col = f" [{it['column'][:40]}]" if it['column'] else ''
            ea = f"  exp={it['expected'][:40]} act={it['actual'][:90]}" if (it['expected'] or it['actual']) else ''
            print(f"{it['severity']:5} {it['layer']:7} {it['check_id']:6} {it['count']:>8,}  {it['check'][:80]}{col}{ea}")
    return F


def _rto(path):
    try:
        xl = pd.ExcelFile(path)
        for s in ['Data', 'RTOSRWithoutChannel', 'RTO', 'Sheet1']:
            if s in xl.sheet_names:
                d = pd.read_excel(path, sheet_name=s)
                if 'UW Budget Cluster' in d.columns and 'Code' in d.columns:
                    d = d.dropna(subset=['UW Budget Cluster'])
                    return {str(k).strip(): ','.join(sorted(set(map(str, g['Code'].dropna().astype(str).str.strip()))))
                            for k, g in d.groupby(d['UW Budget Cluster'].astype(str).str.strip())}
    except Exception:
        return None
    return None


def catalog():
    return {'check_types': CHECK_TYPES, 'ops': OPS, 'no_value_ops': sorted(NO_VALUE_OPS),
            'policies': ['expect', 'optional', 'skip'], 'severities': [ERROR, WARN, INFO]}


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--csv', required=True); ap.add_argument('--profile', required=True)
    for a in ('source', 'rto', 'reference', 'config', 'report'): ap.add_argument('--' + a)
    a = ap.parse_args()
    try: F = run(a.csv, a.profile, a.source, a.rto, a.reference, a.config, a.report)
    except Exception: traceback.print_exc(); sys.exit(2)
    sys.exit(1 if F.has_errors() else 0)
